import html
import re
from io import BytesIO
from pathlib import Path
from urllib.parse import quote_plus
from urllib.request import Request, urlopen
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

BASE = Path(r"c:\proj\itm\OPTATIVA_II")
OUT = BASE / "presentacion_video_1_slides_1_15_diagramas_conceptuales.pptx"
IMG = BASE / "assets_imagenes_videos_1_20"

slides = [
    ("Una decision puede convertirse en un riesgo real", "La historia del Boeing 737 MAX muestra que la arquitectura de software no es un ejercicio abstracto. Cuando un sistema automatizado influye en una operacion critica, una decision de diseno puede afectar la seguridad, la confianza y la vida de las personas.\n\nPor eso, hablar de arquitectura implica hablar tambien de responsabilidad: que supuestos aceptamos, que informacion usamos y que mecanismos de proteccion decidimos incluir."),
    ("El caso Boeing 737 MAX: cuando un supuesto domina el sistema", "En los accidentes de 2018 y 2019, el sistema MCAS podia recibir informacion de un unico sensor de angulo de ataque. Al no contar con una redundancia suficiente ni con una validacion adecuada entre fuentes, una lectura incorrecta podia activar una respuesta automatica peligrosa.\n\nLa leccion arquitectonica no consiste en reducir la tragedia a un error de programacion. El problema estuvo relacionado con decisiones de sistema: dependencia de una fuente unica, supuestos no suficientemente cuestionados y controles de seguridad insuficientes."),
    ("No fue solamente un bug", "Un bug suele describirse como un comportamiento no intencional del software. En este caso, la reflexion es mas profunda: existian decisiones conscientes sobre como integrar el sistema, cuanto automatizar y que nivel de redundancia incorporar.\n\nCuando una arquitectura prioriza rapidez, costo o compatibilidad comercial sobre seguridad y evidencia tecnica, el riesgo no desaparece. Simplemente queda oculto hasta que las condiciones reales lo hacen visible."),
    ("La arquitectura expresa prioridades", "Toda arquitectura es una forma de ordenar prioridades. Si se privilegia reducir tiempos de entrega, se puede ganar velocidad inicial. Si se privilegia minimizar infraestructura, se puede reducir gasto. Pero cada beneficio tiene consecuencias que deben hacerse visibles.\n\nEn sistemas criticos, la pregunta no es solo cuanto cuesta incorporar una proteccion, sino cuanto cuesta operar sin ella. La arquitectura responsable compara ambos lados antes de decidir."),
    ("Una sola fuente de informacion puede convertirse en un punto critico", "Cuando una funcion importante depende de un unico sensor, servicio o repositorio de datos, ese elemento se transforma en un punto unico de falla. Si su informacion es incorrecta, el sistema puede tomar una decision incorrecta con apariencia de certeza.\n\nLa redundancia no significa duplicar todo sin criterio. Significa contar con fuentes independientes, validar discrepancias y definir que debe hacer el sistema cuando la informacion no es confiable."),
    ("La automatizacion necesita limites", "Automatizar una decision no elimina la responsabilidad de disenar sus condiciones de seguridad. Un sistema automatico debe saber cuando actuar, cuando pedir confirmacion y cuando ceder el control a una persona.\n\nLa autonomia sin limites puede hacer que un error pequeno se convierta en una secuencia dificil de detener. Por eso, cada automatizacion necesita reglas de activacion, mecanismos de anulacion y comportamiento seguro ante incertidumbre."),
    ("Las decisiones arquitectonicas tienen consecuencias", "Una eleccion de arquitectura puede afectar escalabilidad, seguridad, privacidad, accesibilidad, disponibilidad y cumplimiento. La consecuencia puede ser tecnica, economica, social o etica.\n\nEsto aplica a cualquier sector: aviacion, salud, banca, educacion o una startup. El contexto cambia, pero la responsabilidad permanece: construir sistemas cuyo comportamiento sea comprensible y controlable."),
    ("El arquitecto no se limita a dibujar", "El trabajo arquitectonico incluye abstraer complejidad, cuestionar supuestos, reconocer riesgos y negociar con las personas que tienen intereses sobre el sistema.\n\nUn diagrama puede ayudar a conversar, pero no reemplaza la investigacion ni la toma de decisiones. La arquitectura se demuestra cuando el equipo puede explicar por que eligio una alternativa y que consecuencias esta dispuesto a aceptar."),
    ("Pensar en los stakeholders cambia la decision", "Los stakeholders son las personas, equipos y organizaciones afectadas por el sistema o capaces de influir en el. Sus expectativas pueden ser diferentes: el negocio busca valor, operaciones busca estabilidad, seguridad busca reducir amenazas y los usuarios buscan una experiencia confiable.\n\nUna decision arquitectonica completa considera esas perspectivas y hace explicitos los conflictos en lugar de ocultarlos bajo una solucion tecnica."),
    ("Comunicar la arquitectura evita perder el razonamiento", "Una arquitectura que solo vive en la memoria de una persona se vuelve dificil de mantener. Cuando esa persona cambia de equipo, las decisiones pierden contexto y el sistema comienza a evolucionar por costumbre o improvisacion.\n\nDocumentar arquitectura significa conservar el razonamiento: que problema existia, que alternativas se evaluaron, que se eligio y que riesgos quedaron abiertos."),
    ("ARCHITECTURE.md como punto de entrada", "Una practica sencilla es crear un archivo ARCHITECTURE.md en la raiz del repositorio. No debe ser una enciclopedia ni repetir cada detalle del codigo. Debe permitir que una persona nueva entienda rapidamente como esta organizado el sistema.\n\nEl documento funciona como una puerta de entrada para explorar el repositorio y como memoria compartida de las decisiones importantes."),
    ("Que debe contener ARCHITECTURE.md", "El documento debe comenzar explicando el proposito general del sistema y el problema que resuelve. Luego debe presentar un mapa breve de modulos, responsabilidades principales y relaciones relevantes.\n\nTambien debe registrar conceptos clave, restricciones conocidas, dependencias externas y riesgos que el equipo todavia debe vigilar. El objetivo es orientar, no reemplazar el codigo ni la documentacion detallada de cada componente."),
    ("Un diagrama sencillo es suficiente cuando responde preguntas", "El diagrama que acompana ARCHITECTURE.md debe ayudar a responder preguntas concretas: que partes existen, como se comunican y donde estan los limites.\n\nNo es necesario dibujar toda la plataforma. Un buen diagrama de contexto o de componentes muestra la estructura suficiente para comprender decisiones importantes sin convertir la arquitectura en un plano imposible de mantener."),
    ("De la documentacion a la accion", "Documentar no es el punto final. La informacion escrita debe ayudar a revisar supuestos, detectar riesgos y tomar mejores decisiones en el siguiente cambio.\n\nCada modificacion importante del sistema deberia preguntarse si cambia una responsabilidad, una dependencia, un contrato o un riesgo. Si la respuesta es afirmativa, la documentacion arquitectonica tambien debe evolucionar."),
    ("La perspectiva arquitectonica amplia el desarrollo", "El desarrollo resuelve funcionalidades concretas; la arquitectura observa como esas funcionalidades conviviran, cambiaran y operaran en el tiempo. Esa perspectiva no compite con programar: ayuda a programar con mayor claridad.\n\nEl objetivo es diseñar sistemas seguros, escalables y efectivos, pero tambien comprensibles, mantenibles y honestos frente a sus limitaciones."),
]

image_queries = [
    "Boeing 737 MAX airplane aviation",
    "airplane flight control cockpit",
    "aircraft sensor aviation safety",
    "risk management engineering decision",
    "redundant sensors aviation system",
    "aircraft automation cockpit controls",
    "aviation safety engineering analysis",
    "software architect whiteboard team",
    "business stakeholders technology meeting",
    "software documentation repository",
    "GitHub repository architecture documentation",
    "software architecture modules diagram",
    "system architecture diagram components",
    "software documentation maintenance",
    "software engineering architecture team",
]

image_explanations = [
    "El avion representa un sistema complejo y critico: una decision de software puede afectar una operacion fisica y la seguridad de sus usuarios.",
    "La cabina y los controles muestran que el software automatizado participa en decisiones operativas que deben tener limites y supervision.",
    "El sensor representa la fuente de datos que alimenta una decision automatica; si esa fuente falla, el sistema puede actuar sobre una premisa equivocada.",
    "El tablero de riesgos representa la necesidad de comparar rapidez y costo contra consecuencias de seguridad antes de aprobar una arquitectura.",
    "Los componentes duplicados representan redundancia: una funcion critica no deberia depender de un unico dato o punto de falla.",
    "Los controles de la cabina representan los limites, alertas y mecanismos de anulacion que una automatizacion responsable debe ofrecer.",
    "La imagen de ingenieria representa que las decisiones arquitectonicas deben evaluarse por sus efectos tecnicos, operativos, economicos y eticos.",
    "La pizarra y el equipo representan al arquitecto facilitando conversaciones entre personas con intereses diferentes sobre el mismo sistema.",
    "La reunion representa a los stakeholders: negocio, operaciones, seguridad y desarrollo deben construir criterios compartidos para decidir.",
    "La documentacion representa la memoria de las decisiones; sin ella, el equipo pierde el razonamiento cuando el sistema evoluciona.",
    "El repositorio representa el lugar donde ARCHITECTURE.md debe vivir junto al codigo para que la estructura del sistema sea facil de consultar.",
    "Los bloques representan modulos y responsabilidades; el documento debe mostrar esta estructura sin describir cada detalle de implementacion.",
    "El diagrama representa una frontera util: muestra componentes y relaciones suficientes para comprender el sistema sin intentar dibujar toda la plataforma.",
    "La imagen representa mantenimiento continuo: cuando una decision cambia, la documentacion debe actualizarse para conservar el contexto.",
    "El equipo representa la perspectiva amplia del desarrollo: construir software tambien implica comprender evolucion, riesgos, usuarios y consecuencias.",
]


def download_images():
    image_dir = BASE / "assets_video_1_boeing_architecture"
    image_dir.mkdir(exist_ok=True)
    paths = []
    used = set()
    for index, search_terms in enumerate(image_queries, start=1):
        output = image_dir / f"slide_{index:02d}.jpg"
        if output.exists():
            output.unlink()
        try:
            url = f"https://www.bing.com/images/search?q={quote_plus(search_terms)}"
            request = Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urlopen(request, timeout=30) as response:
                page = response.read().decode("utf-8", errors="ignore")
            candidates = [html.unescape(x).replace("\\u002f", "/") for x in re.findall(r"murl&quot;:&quot;(https?://[^&]+)", page)]
            for candidate in candidates:
                if candidate in used:
                    continue
                try:
                    image_request = Request(candidate, headers={"User-Agent": "Mozilla/5.0"})
                    with urlopen(image_request, timeout=20) as image_response:
                        picture = Image.open(BytesIO(image_response.read())).convert("RGB")
                    picture.save(output, format="JPEG", quality=88)
                    used.add(candidate)
                    break
                except Exception:
                    continue
        except Exception:
            pass
        paths.append(output if output.exists() and output.stat().st_size > 1000 else None)
    return paths


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(247, 249, 252); shape.line.fill.background()


def add_header(slide, title):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.0))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(16, 51, 97); shape.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.18), Inches(12.2), Inches(0.65))
    p = box.text_frame.paragraphs[0]; p.text = title; p.font.size = Pt(25); p.font.bold = True; p.font.color.rgb = RGBColor(255,255,255)


def add_content(slide, title, text, image_path=None, image_explanation=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.35), Inches(7.55), Inches(5.75))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(255,255,255); card.line.color.rgb = RGBColor(195,208,225)
    h = slide.shapes.add_textbox(Inches(0.95), Inches(1.7), Inches(6.9), Inches(0.7))
    p = h.text_frame.paragraphs[0]; p.text = title; p.font.size = Pt(21); p.font.bold = True; p.font.color.rgb = RGBColor(16,51,97)
    b = slide.shapes.add_textbox(Inches(0.95), Inches(2.45), Inches(6.9), Inches(4.2))
    p = b.text_frame.paragraphs[0]; p.text = text; p.font.size = Pt(17); p.font.color.rgb = RGBColor(38,44,52)
    add_concept_diagram(slide, image_explanation or "La representacion visual conecta los elementos del concepto.", title)


def add_concept_diagram(slide, explanation, title):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(1.35), Inches(4.3), Inches(5.75))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(238, 244, 250); panel.line.color.rgb = RGBColor(174, 195, 218)

    label = slide.shapes.add_textbox(Inches(8.75), Inches(1.65), Inches(3.7), Inches(0.45)).text_frame
    label.clear(); p = label.paragraphs[0]; p.text = "Relacion directa con el concepto"; p.font.bold = True; p.font.size = Pt(15); p.font.color.rgb = RGBColor(16, 51, 97)

    if "avion" in explanation or "sistema complejo" in explanation:
        visual = [("AVION / SISTEMA CRITICO", 9.0, 2.35, RGBColor(207,226,243)), ("SOFTWARE DE CONTROL", 9.0, 3.35, RGBColor(255,235,190)), ("SEGURIDAD DEL USUARIO", 9.0, 4.35, RGBColor(218,242,225))]
    elif "sensor" in explanation or "fuente de datos" in explanation:
        visual = [("SENSOR UNICO", 8.85, 2.35, RGBColor(255,216,216)), ("DECISION AUTOMATICA", 10.45, 3.35, RGBColor(255,235,190)), ("REDUNDANCIA + VALIDACION", 8.85, 4.55, RGBColor(218,242,225))]
    elif "automatizacion" in explanation or "controles" in explanation:
        visual = [("DATO", 8.85, 2.35, RGBColor(207,226,243)), ("AUTOMATISMO", 10.45, 3.35, RGBColor(255,235,190)), ("ALERTA / ANULACION", 8.85, 4.55, RGBColor(218,242,225))]
    elif "stakeholders" in explanation or "equipo" in explanation:
        visual = [("NEGOCIO", 8.75, 2.25, RGBColor(255,235,190)), ("DECISION", 10.45, 3.35, RGBColor(207,226,243)), ("OPERACION + SEGURIDAD", 8.75, 4.55, RGBColor(218,242,225))]
    elif "repositorio" in explanation or "documentacion" in explanation:
        visual = [("REPOSITORIO", 8.85, 2.25, RGBColor(207,226,243)), ("ARCHITECTURE.md", 8.85, 3.35, RGBColor(218,242,225)), ("CODIGO + DECISIONES", 8.85, 4.45, RGBColor(235,235,235))]
    elif "modulos" in explanation or "bloques" in explanation:
        visual = [("PEDIDOS", 8.75, 2.25, RGBColor(207,226,243)), ("RUTEO", 10.35, 2.25, RGBColor(207,226,243)), ("INCIDENCIAS", 9.55, 3.55, RGBColor(218,242,225))]
    elif "diagrama" in explanation or "frontera" in explanation:
        visual = [("ACTOR", 8.75, 2.25, RGBColor(255,235,190)), ("SISTEMA", 10.35, 2.25, RGBColor(207,226,243)), ("SISTEMA EXTERNO", 9.35, 4.0, RGBColor(218,242,225))]
    else:
        visual = [("PROBLEMA", 8.75, 2.25, RGBColor(255,216,216)), ("DECISION", 10.35, 3.35, RGBColor(255,235,190)), ("CONSECUENCIA", 8.75, 4.55, RGBColor(218,242,225))]

    for label_text, x, y, fill in visual:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.0), Inches(0.68))
        box.fill.solid(); box.fill.fore_color.rgb = fill; box.line.color.rgb = RGBColor(103, 124, 148)
        tf = box.text_frame; tf.clear(); p = tf.paragraphs[0]; p.text = label_text; p.alignment = 1; p.font.bold = True; p.font.size = Pt(11); p.font.color.rgb = RGBColor(30, 42, 55)

    caption = slide.shapes.add_textbox(Inches(8.75), Inches(5.55), Inches(3.7), Inches(1.05)).text_frame
    caption.clear(); p = caption.paragraphs[0]; p.text = explanation; p.font.size = Pt(11); p.font.color.rgb = RGBColor(45, 52, 60)


def build():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    for index, (title, text) in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide); add_header(slide, title)
        add_content(slide, title, text, None, image_explanations[index])
    prs.save(OUT)


if __name__ == "__main__":
    build()
