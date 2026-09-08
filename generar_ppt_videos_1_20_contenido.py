from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = r"c:\proj\itm\OPTATIVA_II\presentacion_videos_1_20_contenido.pptx"

videos = [
    {
        "n": 1,
        "title": "Bienvenida: como pensar arquitectura en este curso",
        "explicar": [
            "La arquitectura de software es un sistema de decisiones, no solo diagramas.",
            "El objetivo del curso es aprender a decidir con evidencia tecnica y de negocio.",
            "Trabajaremos un caso unico: plataforma logistica compleja de retail digital.",
        ],
        "apoyo": [
            "Definicion de arquitectura: estructura, responsabilidades y evolucion.",
            "Resultado final esperado: propuesta arquitectonica defendible.",
            "Ejes del curso: problema, requisitos, estructura, diseno, calidad y evolucion.",
        ],
        "ejemplo": "Mapa general del sistema: pedidos, inventario, rutas, repartidores, incidencias y soporte.",
    },
    {
        "n": 2,
        "title": "La arquitectura como conjunto de decisiones",
        "explicar": [
            "Una decision arquitectonica impacta costo, calidad, tiempo y riesgo.",
            "Las decisiones importantes son las que condicionan cambios futuros.",
            "No decidir tambien es una decision con consecuencias.",
        ],
        "apoyo": [
            "Decision = contexto + alternativas + eleccion + consecuencias.",
            "Ejemplos: estilo estructural, limites de modulos, estrategia de integracion.",
            "Herramienta sugerida: ADR (Architecture Decision Record).",
        ],
        "ejemplo": "Comparar monolito modular inicial vs servicios tempranos para pedidos y ruteo.",
    },
    {
        "n": 3,
        "title": "El costo invisible de elegir mal",
        "explicar": [
            "La mala arquitectura no falla al inicio, falla cuando el sistema crece.",
            "El costo oculto aparece en retrabajo, incidentes y lentitud del equipo.",
            "La deuda tecnica se acumula cuando se posterga la calidad estructural.",
        ],
        "apoyo": [
            "Costos invisibles: soporte, correcciones urgentes, caidas, sobretiempo.",
            "Sena les de alerta: cambios pequenos con alto impacto colateral.",
            "Medir costo de cambio por funcionalidad critica.",
        ],
        "ejemplo": "Cambio en asignacion de rutas afecta pagos y notificaciones por acoplamiento excesivo.",
    },
    {
        "n": 4,
        "title": "El arquitecto y las conversaciones dificiles",
        "explicar": [
            "El arquitecto facilita decisiones entre negocio, desarrollo y operaciones.",
            "Debe traducir objetivos de negocio a limites tecnicos comprensibles.",
            "Debe negociar trade-offs sin perder foco en valor.",
        ],
        "apoyo": [
            "Actores: producto, desarrollo, QA, seguridad, operaciones, soporte.",
            "Conflictos tipicos: velocidad vs calidad, costo vs resiliencia.",
            "Tecnica clave: decisiones explicitas con criterios comunes.",
        ],
        "ejemplo": "Definir nivel de disponibilidad objetivo para ventana de alta demanda.",
    },
    {
        "n": 5,
        "title": "Problema esencial vs detalle incidental",
        "explicar": [
            "Problema esencial: lo que realmente duele al negocio y al usuario.",
            "Detalle incidental: herramienta, framework o preferencia tecnica.",
            "Confundir ambos lleva a soluciones complejas para problemas equivocados.",
        ],
        "apoyo": [
            "Pregunta guia: esta decision cambia el valor al cliente?",
            "Esencial primero, tecnologia despues.",
            "Evitar sobreingenieria temprana.",
        ],
        "ejemplo": "Esencial: reducir retrasos de entrega; incidental: elegir lenguaje A o B.",
    },
    {
        "n": 6,
        "title": "Deuda tecnica y acoplamiento: senales tempranas",
        "explicar": [
            "Acoplamiento alto vuelve riesgoso cualquier cambio.",
            "La deuda tecnica debe gestionarse como riesgo de negocio.",
            "Detectar temprano evita refactors costosos despues.",
        ],
        "apoyo": [
            "Indicadores: codigo duplicado, modulos gigantes, dependencias ciclicas.",
            "Cohesion alta y acoplamiento bajo como criterio de salud.",
            "Priorizacion por impacto y probabilidad.",
        ],
        "ejemplo": "Modulo de pedidos conoce detalles de persistencia, ruteo y notificaciones al mismo tiempo.",
    },
    {
        "n": 7,
        "title": "Del problema a las fronteras del sistema",
        "explicar": [
            "Toda arquitectura comienza delimitando fronteras claras.",
            "Las fronteras definen responsabilidades y contratos.",
            "Sin fronteras no hay ownership ni control de complejidad.",
        ],
        "apoyo": [
            "Contexto: que esta dentro del sistema y que esta fuera.",
            "Entradas, salidas y dependencias externas.",
            "Responsables por cada frontera.",
        ],
        "ejemplo": "Fronteras entre sistema central y proveedores de pagos, mapas y ERP.",
    },
    {
        "n": 8,
        "title": "Requisitos que cambian la arquitectura",
        "explicar": [
            "No todos los requisitos afectan decisiones estructurales.",
            "Los requisitos de alto impacto deben priorizarse primero.",
            "La priorizacion debe considerar valor, riesgo y dependencia.",
        ],
        "apoyo": [
            "Requisitos funcionales criticos vs funcionalidad cosm etica.",
            "Relaciones entre requisitos y atributos de calidad.",
            "Backlog priorizado con criterio explicito.",
        ],
        "ejemplo": "Reasignacion automatica de repartidor ante incidente en menos de 30 segundos.",
    },
    {
        "n": 9,
        "title": "Calidad medible: rendimiento, seguridad y cambio",
        "explicar": [
            "La calidad debe expresarse como escenarios medibles.",
            "Rendimiento, seguridad y mantenibilidad condicionan la arquitectura.",
            "Si no se mide, no se puede defender ni mejorar.",
        ],
        "apoyo": [
            "Plantilla de escenario: contexto, estimulo, respuesta, metrica.",
            "Atributos base: latencia, disponibilidad, seguridad, modificabilidad.",
            "Uso de escenarios para comparar alternativas.",
        ],
        "ejemplo": "95% de solicitudes de creacion de pedido con latencia menor a 2 segundos en pico.",
    },
    {
        "n": 10,
        "title": "Costo total de operacion y costo de cambio",
        "explicar": [
            "La arquitectura debe evaluarse por costo total de vida.",
            "Construir es solo una parte del costo real.",
            "El costo de cambio define sostenibilidad del producto.",
        ],
        "apoyo": [
            "TCO: desarrollo, infraestructura, monitoreo, soporte y mejora.",
            "Costo de cambio: esfuerzo para adaptar reglas y flujos.",
            "Supuestos economicos que deben declararse.",
        ],
        "ejemplo": "Comparar costo de operar un monolito modular frente a servicios por dominio durante 12 meses.",
    },
    {
        "n": 11,
        "title": "Disenar para la incertidumbre",
        "explicar": [
            "La arquitectura debe aceptar que no todo se conoce al inicio.",
            "Se deben favorecer decisiones reversibles en fases tempranas.",
            "Aprender rapido reduce costo de equivocarse.",
        ],
        "apoyo": [
            "Supuestos explicitos con fecha de validacion.",
            "Experimentos tecnicos pequenos antes de escalar.",
            "Plan de reevaluacion de decisiones criticas.",
        ],
        "ejemplo": "Piloto de eventos en un flujo de notificaciones antes de extender al sistema completo.",
    },
    {
        "n": 12,
        "title": "Elegir estructura sin seguir modas",
        "explicar": [
            "La estructura se elige por contexto, no por tendencia.",
            "Cada estilo tiene beneficios y costos concretos.",
            "La comparacion debe hacerse con criterios objetivos.",
        ],
        "apoyo": [
            "Opciones comunes: capas, monolito modular, servicios.",
            "Criterios: equipo, costo, riesgo, rendimiento, evolucion.",
            "Decision trazable con matriz comparativa.",
        ],
        "ejemplo": "Seleccion estructural inicial para volumen medio y equipo pequeno.",
    },
    {
        "n": 13,
        "title": "Cliente-servidor y capas que si protegen",
        "explicar": [
            "Las capas son utiles si cada una tiene responsabilidad clara.",
            "Mezclar dominio con infraestructura rompe mantenibilidad.",
            "La direccion de dependencias debe ser intencional.",
        ],
        "apoyo": [
            "Presentacion, aplicacion, dominio, infraestructura.",
            "Regla: lo externo no domina lo interno.",
            "Flujos de peticion bien delimitados.",
        ],
        "ejemplo": "Flujo crear pedido sin logica de negocio en el controlador HTTP.",
    },
    {
        "n": 14,
        "title": "Monolito modular: empezar pequeno con futuro",
        "explicar": [
            "El monolito modular puede ser una decision estrategica inicial.",
            "Modularidad temprana prepara el camino para crecer.",
            "No se trata de tamano, se trata de limites internos.",
        ],
        "apoyo": [
            "Modulos por dominio: pedidos, inventario, ruteo, incidencias.",
            "Contratos internos entre modulos.",
            "Pruebas de modulo para evitar acoplamiento accidental.",
        ],
        "ejemplo": "Separar logica de asignacion de rutas del flujo de pagos y notificaciones.",
    },
    {
        "n": 15,
        "title": "Servicios y contratos evolutivos",
        "explicar": [
            "Los contratos estables permiten evolucionar sin romper consumidores.",
            "El versionado evita cambios disruptivos.",
            "El contrato debe expresar intencion de negocio.",
        ],
        "apoyo": [
            "Entradas, salidas, errores y semantica del contrato.",
            "Compatibilidad hacia atras como regla de evolucion.",
            "Documentacion como parte del contrato.",
        ],
        "ejemplo": "Contrato de servicio para asignar repartidor con codigos de error estandar.",
    },
    {
        "n": 16,
        "title": "Eventos, consistencia y fallos parciales",
        "explicar": [
            "En sistemas distribuidos la consistencia total inmediata no siempre es viable.",
            "La consistencia eventual requiere controles claros.",
            "Los fallos parciales se disenan, no se ignoran.",
        ],
        "apoyo": [
            "Conceptos: idempotencia, reintento, compensacion.",
            "Dise no de flujos asincronos con trazabilidad.",
            "Tratamiento de duplicados y orden de eventos.",
        ],
        "ejemplo": "Evento PedidoCreado activa inventario y ruteo; si falla ruteo se ejecuta compensacion.",
    },
    {
        "n": 17,
        "title": "Microservicios: cuando si y cuando no",
        "explicar": [
            "Microservicios no son obligatorios para toda arquitectura.",
            "Separar servicios implica costo operativo alto.",
            "La decision depende de contexto, escala y madurez del equipo.",
        ],
        "apoyo": [
            "Prerequisitos: observabilidad, automatizacion, gobierno de contratos.",
            "Riesgos: latencia, consistencia, complejidad de despliegue.",
            "Alternativa valida: monolito modular bien disenado.",
        ],
        "ejemplo": "Mantener monolito modular hasta que ruteo y pedidos requieran escalamiento independiente.",
    },
    {
        "n": 18,
        "title": "SOLID para reducir acoplamiento",
        "explicar": [
            "SOLID mejora legibilidad, extensibilidad y testabilidad.",
            "No es teoria aislada: influye en la arquitectura diaria.",
            "Aplicarlo evita clases gigantes y dependencias fragiles.",
        ],
        "apoyo": [
            "Responsabilidad unica, abierto/cerrado, inversion de dependencias.",
            "Enfoque pragmatico: aplicar donde reduce dolor real.",
            "Refactor pequeno con impacto medible.",
        ],
        "ejemplo": "Separar calculo de ruta, validacion de pedido y envio de notificacion en componentes independientes.",
    },
    {
        "n": 19,
        "title": "Direccion de dependencias y limites del codigo",
        "explicar": [
            "El dominio debe estar protegido de tecnologia cambiante.",
            "Las dependencias deben apuntar hacia reglas de negocio.",
            "Los puertos y adaptadores sostienen ese aislamiento.",
        ],
        "apoyo": [
            "Dependencia permitida: infraestructura depende de dominio.",
            "Dependencia prohibida: dominio depende de framework o BD.",
            "Diseno de interfaces para servicios externos.",
        ],
        "ejemplo": "Adaptador de pagos implementa puerto de cobro sin contaminar el modelo de dominio.",
    },
    {
        "n": 20,
        "title": "Patrones que resuelven problemas reales",
        "explicar": [
            "Los patrones son herramientas, no metas.",
            "Cada patron debe justificar su costo de complejidad.",
            "Elegir patron correcto depende del problema concreto.",
        ],
        "apoyo": [
            "Patrones utiles en este contexto: Adapter, Strategy, Facade.",
            "Criterio de seleccion: problema, alternativa, trade-off.",
            "Evaluar costo de retiro si deja de ser necesario.",
        ],
        "ejemplo": "Usar Strategy para politicas de asignacion y Adapter para integracion con proveedor de mapas.",
    },
]


def paint_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(245, 248, 252)
    bg.line.fill.background()


def add_header(slide, title, subtitle):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(1.0))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(21, 57, 104)
    band.line.fill.background()

    t = slide.shapes.add_textbox(Inches(0.45), Inches(0.16), Inches(12.3), Inches(0.68)).text_frame
    t.clear()
    p = t.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(25)
    p.font.color.rgb = RGBColor(255, 255, 255)

    st = slide.shapes.add_textbox(Inches(0.65), Inches(1.2), Inches(12.0), Inches(0.5)).text_frame
    st.clear()
    sp = st.paragraphs[0]
    sp.text = subtitle
    sp.font.bold = True
    sp.font.size = Pt(17)
    sp.font.color.rgb = RGBColor(21, 57, 104)


def add_card(slide, x, y, w, h, title, lines, title_color=RGBColor(21, 57, 104)):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = RGBColor(193, 206, 224)

    tx = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.18), Inches(w - 0.5), Inches(h - 0.35)).text_frame
    tx.clear()

    tp = tx.paragraphs[0]
    tp.text = title
    tp.font.bold = True
    tp.font.size = Pt(15)
    tp.font.color.rgb = title_color

    for line in lines:
        p = tx.add_paragraph()
        p.text = "- " + line
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(40, 46, 54)


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint_background(slide)
    add_header(slide, "Presentacion de Contenidos | Videos 1-20", "Arquitectura de Software")

    add_card(
        slide,
        0.8,
        2.0,
        11.7,
        3.7,
        "Proposito de esta presentacion",
        [
            "Servir como material de apoyo para exponer cada video con profundidad.",
            "Incluir contenido directo para lectura y explicacion en clase/video.",
            "Mantener foco en el caso oficial: plataforma logistica compleja de retail digital.",
        ],
    )


def add_video_slide(prs, v):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint_background(slide)
    add_header(slide, f"Video {v['n']}: {v['title']}", "Contenido para exponer")

    add_card(slide, 0.6, 1.85, 6.1, 2.55, "Que explicar", v["explicar"])
    add_card(slide, 6.8, 1.85, 6.0, 2.55, "Material de apoyo", v["apoyo"])
    add_card(slide, 0.6, 4.6, 12.2, 2.15, "Ejemplo aplicado al caso logistica retail", [v["ejemplo"]])


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    add_cover(prs)
    for v in videos:
        add_video_slide(prs, v)

    prs.save(OUT)


if __name__ == "__main__":
    main()
