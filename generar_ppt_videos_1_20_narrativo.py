import html
from io import BytesIO
import re
from pathlib import Path
from urllib.parse import quote_plus, urlparse
from urllib.request import Request, urlopen

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

BASE_DIR = Path(r"c:\proj\itm\OPTATIVA_II")
OUT = BASE_DIR / "presentacion_videos_1_20_narrativa_visual_v3.pptx"
IMG_DIR = BASE_DIR / "assets_imagenes_videos_1_20"

TOPICS = [
    {
        "title": "Arquitectura de software: decisiones, no solo diagramas",
        "queries": ["software architecture whiteboard team", "system architecture blueprint technology", "retail logistics technology network"],
        "p1": "Durante mucho tiempo se asumio que hacer arquitectura era producir una gran coleccion de diagramas y documentos antes de comenzar a construir. Ese enfoque parecia ordenado, pero en la practica dejaba a los equipos con una falsa sensacion de control: habia mucho material visual y poca claridad sobre que decisiones realmente sostenian el negocio.",
        "p2": "Una arquitectura madura no compite con el desarrollo: lo habilita. Define limites, responsabilidades, contratos y reglas de evolucion para que el software soporte cambios reales sin colapsar. Los diagramas siguen siendo utiles, pero solo cuando nacen de decisiones explicitas y comprobables.",
        "p3": "La idea central de este recorrido es simple: cada vez que hablamos de arquitectura, en realidad estamos hablando de decisiones con consecuencias tecnicas, economicas y operativas. Esa perspectiva evita caer en soluciones vistosas pero fragiles.",
        "case": "En una plataforma logistica de retail digital, la pregunta clave no es cuan bonito luce el diagrama. La pregunta clave es si una variacion de demanda, una falla de un proveedor o un cambio en reglas de entrega puede resolverse sin romper todo el sistema.",
    },
    {
        "title": "Decidir con evidencia tecnica y de negocio",
        "queries": ["business technology strategy meeting", "software metrics analytics dashboard", "logistics business decision meeting"],
        "p1": "Una decision arquitectonica no deberia surgir de preferencias personales ni de modas tecnicas. Deberia surgir de evidencia. Por evidencia de negocio entendemos objetivos, restricciones economicas, perfil de usuarios, estacionalidad de demanda y riesgos de operacion.",
        "p2": "Por evidencia tecnica entendemos estado actual del sistema, deuda acumulada, limitaciones de talento, costo de integraciones, observabilidad disponible y capacidad real de despliegue. Sin esa doble evidencia, la arquitectura queda reducida a opinion.",
        "p3": "Cuando negocio y tecnica se leen juntas, las conversaciones cambian: ya no se discute quien tiene la razon, se discute que alternativa tiene mejor relacion entre valor, costo y riesgo en el contexto real.",
        "case": "Si el negocio exige reducir retrasos sin duplicar gasto operativo, la arquitectura debe priorizar trazabilidad, reaccion ante incidencias y control de capacidad antes que complejidad prematura.",
    },
    {
        "title": "Costo invisible y deuda arquitectonica",
        "queries": ["technical debt software code", "software maintenance developer code", "business cost risk technology"],
        "p1": "Muchas malas decisiones no duelen el primer mes. Duelen cuando el producto necesita cambiar con velocidad y cada ajuste tarda semanas. Ese atraso se vuelve gasto: mas incidentes, mas soporte, mas retrabajo, mas horas improductivas.",
        "p2": "La deuda arquitectonica no es solo codigo sucio. Es perdida de opcionalidad. Cada acoplamiento innecesario reduce capacidad de evolucion y eleva el costo de cualquier iniciativa futura.",
        "p3": "Visualizar ese costo invisible es una habilidad estrategica. Un equipo que lo entiende deja de celebrar entregas rapidas pero fragiles y empieza a proteger capacidad de cambio como un activo de negocio.",
        "case": "Si cambiar reglas de ruteo obliga a tocar pagos, inventario y notificaciones, el sistema esta pagando intereses de deuda. La solucion no es mas esfuerzo manual: es redisenar limites y contratos.",
    },
    {
        "title": "Conversaciones dificiles y rol arquitectonico",
        "queries": ["software stakeholder workshop", "team negotiation technology project", "operations security business meeting"],
        "p1": "La arquitectura no se define en aislamiento tecnico. Surge en conversaciones donde existen intereses en tension: negocio pide velocidad, operaciones pide estabilidad, seguridad exige controles, desarrollo busca simplicidad.",
        "p2": "El rol arquitectonico consiste en transformar ese conflicto en decisiones transparentes. No se trata de ganar discusiones, se trata de explicitar criterios: impacto en cliente, riesgo operativo, costo total y sostenibilidad de cambio.",
        "p3": "Cuando la conversacion se vuelve trazable, el equipo deja de depender de autoridad personal y empieza a construir criterio compartido. Esa transicion mejora calidad de decision y reduce friccion organizacional.",
        "case": "En retail digital, aceptar todos los pedidos en pico puede mejorar ventas de corto plazo y destruir experiencia por incumplimiento. La arquitectura debe sostener ese equilibrio con reglas claras.",
    },
    {
        "title": "Problema esencial y complejidad incidental",
        "queries": ["design thinking problem solving board", "software requirements discussion", "technology choice decision team"],
        "p1": "Un error clasico es enamorarse de soluciones antes de formular bien el problema. En arquitectura, esa inversion de orden genera sistemas complejos que optimizan detalles secundarios y descuidan el objetivo principal del producto.",
        "p2": "El problema esencial es aquello que, si no se resuelve, invalida la propuesta de valor. Lo incidental son elecciones de herramienta o implementacion que pueden variar sin afectar la esencia del negocio.",
        "p3": "Separar ambos niveles evita sobredisenar y ayuda a mantener foco. Primero se protege el resultado de negocio; despues se elige la tecnologia que mejor lo habilita con el menor costo de complejidad.",
        "case": "Reducir retrasos y mejorar trazabilidad es esencial. Elegir entre dos frameworks equivalentes es incidental mientras ambos cumplan los objetivos de calidad y cambio.",
    },
    {
        "title": "Acoplamiento, cohesion y salud del sistema",
        "queries": ["software refactoring code developer", "modular software architecture code", "software system dependency graph"],
        "p1": "La salud arquitectonica se percibe en la facilidad de cambio. Si una mejora pequena obliga a modificar demasiadas piezas, existe acoplamiento excesivo. Si cada modulo tiene una responsabilidad clara, la cohesion mejora y el sistema se vuelve maniobrable.",
        "p2": "Acoplamiento alto no solo complica codigo: tambien paraliza decisiones de negocio porque cada iniciativa parece riesgosa. Cohesion alta acelera entregas porque limita el impacto de cambio y facilita pruebas confiables.",
        "p3": "Observar dependencias, ciclos, duplicacion de reglas y fragilidad de despliegue permite detectar deterioro temprano y corregirlo antes de que la deuda se vuelva estructural.",
        "case": "Cuando pedidos conoce detalles de ruteo, pagos y notificaciones en la misma capa, conviene separar responsabilidades para recuperar control evolutivo.",
    },
    {
        "title": "Fronteras y contexto del sistema",
        "queries": ["system context architecture diagram", "software integration external systems", "logistics platform ecosystem"],
        "p1": "Sin fronteras claras, la arquitectura se vuelve ambigua. Definir contexto significa declarar que responsabilidades son propias del sistema y cuales dependen de actores o plataformas externas.",
        "p2": "Ese mapa inicial ordena discusiones tecnicas y operativas. Permite diferenciar fallas internas de fallas de dependencia, definir contratos realistas y planear mecanismos de contingencia.",
        "p3": "Las fronteras bien descritas tambien facilitan gobernanza: ownership de equipos, rutas de escalamiento e impacto de cambios en cada interfaz.",
        "case": "En ultima milla, pagos, mapas y ERP son externos. El sistema central debe protegerse con contratos, reintentos y observabilidad sin asumir control que no posee.",
    },
    {
        "title": "Requisitos que realmente mueven arquitectura",
        "queries": ["software requirements workshop", "product backlog prioritization", "logistics delivery requirements"],
        "p1": "No todo requisito tiene peso arquitectonico. Algunos cambian presentacion; otros exigen decisiones de estructura, datos o integracion. Identificar esa diferencia temprano ahorra meses de retrabajo.",
        "p2": "Priorizar requiere mirar valor de negocio, criticidad operativa, dependencia tecnica y costo de incumplimiento. Esta mezcla evita roadmaps vistosos pero poco estrategicos.",
        "p3": "Cuando la priorizacion es explicita, la arquitectura responde a lo que mas impacta resultado y no a lo que suena mas urgente en la semana.",
        "case": "Reasignar rutas en segundos durante incidentes afecta estructura y observabilidad. Cambiar etiquetas de interfaz no tiene ese mismo peso.",
    },
    {
        "title": "Calidad medible como base de decision",
        "queries": ["software performance monitoring dashboard", "software security testing screen", "system availability monitoring"],
        "p1": "Calidad expresada como adjetivo no sirve para decidir. Rapido, seguro o escalable necesitan definicion cuantitativa: en que escenario, bajo que carga y con que umbral aceptable.",
        "p2": "Los escenarios de calidad convierten aspiraciones en criterios verificables. Gracias a ellos, dos alternativas se comparan por evidencia y no por percepcion.",
        "p3": "Este enfoque vuelve mas objetiva la conversacion con negocio y mejora trazabilidad de compromisos tecnicos en el tiempo.",
        "case": "Si pedidos debe responder en menos de dos segundos con carga pico, la arquitectura debe proteger camino critico y aislar procesos no urgentes.",
    },
    {
        "title": "Costo total de operacion y costo de cambio",
        "queries": ["cloud computing operations cost", "software operations team monitoring", "technology lifecycle planning"],
        "p1": "El costo real de una arquitectura no termina al desplegar la primera version. Incluye soporte, monitoreo, incidentes, evolucion funcional y entrenamiento del equipo.",
        "p2": "Evaluar costo de cambio es crucial: cuanto esfuerzo exige adaptar una regla central sin afectar estabilidad. Un sistema barato de construir puede ser carisimo de mantener.",
        "p3": "La mirada de ciclo de vida evita decisiones de corto plazo que comprometen sostenibilidad financiera y tecnica.",
        "case": "Separar responsabilidades desde el inicio puede aumentar esfuerzo inicial, pero reduce costo de adaptacion cuando cambian reglas de entrega o integraciones externas.",
    },
    {
        "title": "Disenar para incertidumbre",
        "queries": ["software prototype experiment team", "technology risk planning", "incremental product development"],
        "p1": "Toda arquitectura nace con incertidumbre: demanda, comportamiento de usuarios, performance de proveedores y madurez del equipo. Fingir certeza total produce disenos rigidos y caros de corregir.",
        "p2": "La estrategia sensata es identificar supuestos, priorizar decisiones reversibles y ejecutar experimentos pequenos que reduzcan riesgo antes de comprometer grandes inversiones.",
        "p3": "Disenar para incertidumbre no es improvisar; es aprender sistematicamente y ajustar con disciplina.",
        "case": "Validar un flujo asincrono en un componente acotado antes de expandirlo al sistema completo reduce probabilidad de fallas estructurales.",
    },
    {
        "title": "Elegir estructura con criterios y no con modas",
        "queries": ["software architecture planning diagram", "technology architecture comparison", "engineering design decision"],
        "p1": "Cada estilo arquitectonico resuelve ciertos problemas y trae costos concretos. Elegir por tendencia suele ocultar esas concesiones hasta que aparecen en produccion.",
        "p2": "Elegir con criterio implica ponderar contexto: presupuesto, tamano de equipo, horizonte de crecimiento, criticidad operativa y tolerancia a complejidad.",
        "p3": "Cuando esos criterios son publicos, la decision se vuelve defendible ante negocio y tecnica por igual.",
        "case": "Para un equipo pequeno con crecimiento moderado, un monolito modular puede ofrecer mejor relacion valor-complejidad que una separacion temprana en muchos servicios.",
    },
    {
        "title": "Capas que protegen de verdad",
        "queries": ["layered software architecture diagram", "clean code layers programming", "web application request flow"],
        "p1": "Separar por capas no es un ejercicio estetico. Su objetivo es evitar que detalles transitorios contaminen reglas de negocio que deberian permanecer estables.",
        "p2": "Cuando la capa externa domina a la interna, el sistema se vuelve fragil. Cuando la direccion de dependencias protege el nucleo, el cambio se vuelve mas seguro.",
        "p3": "La arquitectura por capas funciona si cada capa tiene frontera real y reglas de acceso verificables.",
        "case": "El caso de uso crear pedido deberia probarse sin servidor HTTP ni base real. Si no es posible, el limite entre capas esta roto.",
    },
    {
        "title": "Monolito modular como punto de partida sano",
        "queries": ["modular software architecture blocks", "software modules code organization", "retail logistics software platform"],
        "p1": "Un monolito modular bien disenado puede ser una estrategia altamente profesional. Permite velocidad inicial sin renunciar a orden interno y prepara una evolucion gradual.",
        "p2": "La clave esta en definir modulos por responsabilidad de negocio, no por conveniencia tecnica del framework. Cada modulo necesita contratos y limites claros.",
        "p3": "Con esta base, el sistema puede escalar organizacionalmente antes de escalar topologicamente.",
        "case": "Pedidos, inventario, ruteo e incidencias pueden convivir en un despliegue unico mientras mantengan interfaces internas estables y pruebas por frontera.",
    },
    {
        "title": "Contratos evolutivos para servicios e integraciones",
        "queries": ["api contract documentation", "software api versioning", "system integration interface"],
        "p1": "Los contratos son acuerdos tecnicos de negocio: que se env ia, que se responde, que errores existen y como evoluciona la interfaz sin romper consumidores.",
        "p2": "Sin disciplina de contrato, cada cambio interno genera friccion externa. Con versionado y compatibilidad, la evolucion se vuelve predecible.",
        "p3": "La calidad de una arquitectura distribuida depende en gran medida de la calidad de sus contratos.",
        "case": "Una regla nueva en asignacion de repartidores no deberia invalidar clientes existentes si el contrato mantiene compatibilidad y semantica estable.",
    },
    {
        "title": "Eventos, consistencia y recuperacion",
        "queries": ["event driven architecture data flow", "distributed systems message queue", "logistics delivery tracking technology"],
        "p1": "Arquitectura orientada a eventos aporta desacoplamiento, pero exige madurez en consistencia y recuperacion. En sistemas reales, los mensajes se duplican, se retrasan o llegan fuera de orden.",
        "p2": "Por eso se necesitan idempotencia, reintentos controlados, trazabilidad y estrategias de compensacion. Sin esos mecanismos, la asincronia se convierte en fuente de errores silenciosos.",
        "p3": "Una buena arquitectura de eventos no promete ausencia de fallos; promete capacidad de detectar, aislar y corregir fallos sin colapsar el negocio.",
        "case": "Si pago confirma y ruteo falla, el sistema debe reintentar o compensar de forma automatizada con trazabilidad completa para soporte.",
    },
    {
        "title": "Microservicios: criterio antes que entusiasmo",
        "queries": ["microservices architecture containers", "kubernetes operations team", "distributed system observability"],
        "p1": "Microservicios pueden aportar autonomia y escalado selectivo, pero tambien elevan costos de observabilidad, despliegue, seguridad y coordinacion de equipos.",
        "p2": "Adoptarlos con baja madurez operativa genera sistemas distribuidos imposibles de gobernar. El resultado suele ser menos velocidad y mas incidentes.",
        "p3": "La decision correcta no es la mas moderna: es la que el contexto puede sostener con calidad.",
        "case": "Si aun no existen metricas claras, pipelines confiables y ownership de dominios, conviene fortalecer base modular antes de fragmentar el sistema.",
    },
    {
        "title": "SOLID como herramienta practica",
        "queries": ["clean code refactoring programming", "software design principles code", "developer testing software architecture"],
        "p1": "SOLID no es una lista para memorizar, es un lente para detectar disenos fragiles. Aplicado con criterio, reduce impacto colateral, mejora testabilidad y acelera mantenimiento.",
        "p2": "Su valor real aparece cuando se conecta con problemas concretos: clases sobrecargadas, extensiones riesgosas o dependencias rigidas a implementaciones externas.",
        "p3": "Adoptarlo de forma pragmatica permite elevar calidad sin caer en ceremonias innecesarias.",
        "case": "Separar validacion de pedido, calculo de ruta y notific acion reduce acoplamiento y permite evolucion independiente de cada politica.",
    },
    {
        "title": "Direccion de dependencias y limites del codigo",
        "queries": ["hexagonal architecture ports adapters", "software dependency inversion diagram", "payment api adapter integration"],
        "p1": "La direccion de dependencias define quien manda en el sistema. Cuando el dominio depende de tecnologia, cada cambio de plataforma impacta reglas de negocio.",
        "p2": "Con puertos y adaptadores, el dominio expresa necesidades y la infraestructura implementa detalles. Ese orden protege el nucleo frente a cambios externos.",
        "p3": "Este principio transforma arquitectura en una red de limites explicitos y evolutivos.",
        "case": "El dominio solicita cobrar pedido por un puerto abstracto; un adaptador concreto integra el proveedor de pagos sin contaminar reglas centrales.",
    },
    {
        "title": "Patrones para problemas reales",
        "queries": ["software design patterns diagram", "adapter design pattern programming", "strategy pattern software design"],
        "p1": "Un patron bien elegido reduce complejidad accidental. Un patron mal elegido agrega capas innecesarias y dificulta mantenimiento. La clave es partir del problema, no de la herramienta.",
        "p2": "Conviene comparar alternativa simple contra alternativa con patron y justificar el costo adicional solo cuando haya beneficio tangible en flexibilidad o estabilidad.",
        "p3": "La madurez arquitectonica no se mide por cantidad de patrones usados, sino por claridad de decisiones y capacidad de evolucion.",
        "case": "Adapter para aislar integraciones externas y Strategy para variar asignacion de rutas tienen sentido si realmente cambian politicas y proveedores con frecuencia.",
    },
]


def ensure_images():
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    paths = {}
    image_number = 1
    used_urls = set()
    for topic_number, t in enumerate(TOPICS, start=1):
        paths[topic_number] = []
        for slide_number, raw_query in enumerate(t["queries"], start=1):
            img_path = IMG_DIR / f"topic_{topic_number:02d}_slide_{slide_number:02d}.jpg"
            if img_path.exists():
                img_path.unlink()
            query = quote_plus(raw_query)
            selected_url = None
            try:
                search_terms = [query, quote_plus("software technology logistics " + str(image_number))]
                candidates = []
                for search_term in search_terms:
                    search_url = f"https://www.bing.com/images/search?q={search_term}"
                    request = Request(search_url, headers={"User-Agent": "Mozilla/5.0"})
                    with urlopen(request, timeout=20) as response:
                        search_html = response.read().decode("utf-8", errors="ignore")
                    candidates.extend(
                        html.unescape(match).replace("\\u002f", "/")
                        for match in re.findall(r"murl&quot;:&quot;(https?://[^&]+)", search_html)
                    )
                for candidate in candidates:
                    if candidate in used_urls:
                        continue
                    try:
                        image_request = Request(
                            candidate,
                            headers={
                                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) ArchitectureCourse/1.0",
                                "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
                            },
                        )
                        with urlopen(image_request, timeout=30) as image_response:
                            source_image = Image.open(BytesIO(image_response.read())).convert("RGB")
                        source_image.save(img_path, format="JPEG", quality=88)
                        selected_url = candidate
                        used_urls.add(selected_url)
                        break
                    except Exception:
                        continue
            except Exception:
                img_path = None
            if img_path and (not img_path.exists() or img_path.stat().st_size < 1000):
                img_path = None
            paths[topic_number].append(img_path)
            image_number += 1
    return paths


def paint_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(246, 249, 253)
    bg.line.fill.background()


def add_header(slide, title):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.92))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(17, 50, 98)
    band.line.fill.background()

    tf = slide.shapes.add_textbox(Inches(0.45), Inches(0.12), Inches(12.5), Inches(0.7)).text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(255, 255, 255)


def add_text_card(slide, heading, paragraphs):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.2), Inches(12.0), Inches(6.0))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = RGBColor(194, 208, 226)

    ht = slide.shapes.add_textbox(Inches(0.95), Inches(1.45), Inches(11.4), Inches(0.7)).text_frame
    ht.clear()
    h = ht.paragraphs[0]
    h.text = heading
    h.font.bold = True
    h.font.size = Pt(20)
    h.font.color.rgb = RGBColor(17, 50, 98)

    body = slide.shapes.add_textbox(Inches(0.95), Inches(2.18), Inches(11.4), Inches(4.8)).text_frame
    body.clear()
    for i, para in enumerate(paragraphs):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = para
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(38, 45, 54)


def add_visual_slide(prs, title, t, image_path):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    paint_bg(s)
    add_header(s, title)

    left = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.2), Inches(6.55), Inches(6.0))
    left.fill.solid()
    left.fill.fore_color.rgb = RGBColor(255, 255, 255)
    left.line.color.rgb = RGBColor(194, 208, 226)

    ltf = s.shapes.add_textbox(Inches(0.85), Inches(1.45), Inches(5.95), Inches(5.4)).text_frame
    ltf.clear()

    p0 = ltf.paragraphs[0]
    p0.text = "Desarrollo aplicado"
    p0.font.bold = True
    p0.font.size = Pt(17)
    p0.font.color.rgb = RGBColor(17, 50, 98)

    for line in [
        t["p2"],
        t["p3"],
        "Caso transversal: " + t["case"],
    ]:
        p = ltf.add_paragraph()
        p.text = line
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(38, 45, 54)

    right = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.2), Inches(5.45), Inches(6.0))
    right.fill.solid()
    right.fill.fore_color.rgb = RGBColor(255, 255, 255)
    right.line.color.rgb = RGBColor(194, 208, 226)

    if image_path and image_path.exists() and image_path.stat().st_size > 0:
        s.shapes.add_picture(str(image_path), Inches(7.55), Inches(1.5), width=Inches(4.95), height=Inches(3.35))
        caption = s.shapes.add_textbox(Inches(7.55), Inches(4.95), Inches(4.95), Inches(2.0)).text_frame
        caption.clear()
        c0 = caption.paragraphs[0]
        c0.text = "Imagen de referencia obtenida desde internet"
        c0.font.bold = True
        c0.font.size = Pt(12)
        c0.font.color.rgb = RGBColor(17, 50, 98)

        c1 = caption.add_paragraph()
        c1.text = t["case"]
        c1.font.size = Pt(12)
        c1.font.color.rgb = RGBColor(38, 45, 54)
    else:
        fallback = s.shapes.add_textbox(Inches(7.55), Inches(1.6), Inches(4.95), Inches(4.8)).text_frame
        fallback.clear()
        f0 = fallback.paragraphs[0]
        f0.text = "No fue posible descargar imagen en este intento."
        f0.font.bold = True
        f0.font.size = Pt(13)
        f0.font.color.rgb = RGBColor(120, 37, 37)


def build():
    image_paths = ensure_images()

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # Cover
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    paint_bg(cover)
    add_header(cover, "Arquitectura de Software - Recorrido narrativo profundo")
    add_text_card(
        cover,
        "Presentacion expandida con apoyo visual",
        [
            "Este material desarrolla los veinte temas iniciales con texto extendido para lectura en pantalla y explicacion continua.",
            "Cada bloque profundiza contexto, implicaciones tecnicas y criterio de decision aplicado al caso de logistica retail.",
            "Se incorporan imagenes relevantes obtenidas desde internet mediante busqueda por palabras clave para reforzar comprension visual.",
        ],
    )

    for idx, t in enumerate(TOPICS, start=1):
        title = t["title"]

        # Slide 1: marco conceptual ampliado
        s1 = prs.slides.add_slide(prs.slide_layouts[6])
        paint_bg(s1)
        add_header(s1, title)
        add_text_card(
            s1,
            "Fundamento y contexto",
            [t["p1"], t["p2"], t["p3"]],
        )

        # Slide 2: desarrollo aplicado + imagen de internet
        add_visual_slide(prs, title, t, image_paths.get(idx, [None, None, None])[1])

        # A third slide closes the idea with a second visual related to the case.
        add_visual_slide(prs, title, t, image_paths.get(idx, [None, None, None])[2])

    prs.save(OUT)


if __name__ == "__main__":
    build()
